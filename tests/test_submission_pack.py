"""Submission pack — offline (guideline pages are mocked).

Run:  python tests/test_submission_pack.py
"""
import io
import os
import sys
import tempfile
import zipfile
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import ingest, profiles, submission_pack as sp  # noqa: E402

GUIDE = (
    "Guide for authors. Submissions must include a cover letter. This journal uses "
    "double anonymized peer review, so the anonymized manuscript must not include author "
    "names. Highlights are mandatory: include 3 to 5 bullet points, maximum 85 characters, "
    "including spaces, per bullet point. A graphical abstract is optional but encouraged. "
    "Authors must declare competing interests. Please suggest 3 potential reviewers. "
    "Figures should be uploaded as separate files in TIFF format, at least 300 dpi. "
    "The manuscript should not exceed 5,000 words.")


def _docx(tmp):
    from docx import Document
    from PIL import Image
    d = Document()
    d.add_paragraph("Deep Learning for Fault Detection in Wind Turbines")
    d.add_paragraph("Asha Rao, Vikram Singh")
    d.add_paragraph("Abstract: Wind turbine faults are costly. We propose a compact model. "
                    "It improves detection accuracy by 12 percent. The model runs in real time.")
    d.add_paragraph("Keywords: fault detection, wind energy, deep learning")
    d.add_heading("Introduction", 1)
    d.add_paragraph("Wind energy is growing quickly and maintenance costs dominate operations "
                    "for many operators across the world today.")
    img = os.path.join(tmp, "f.png")
    Image.new("RGB", (1200, 800), (10, 120, 200)).save(img)
    d.add_picture(img)
    d.add_paragraph("Figure 1. Proposed architecture")
    d.add_heading("Conclusion", 1)
    d.add_paragraph("The model detects faults early. It needs little data.")
    d.add_heading("Acknowledgements", 1)
    d.add_paragraph("We thank the Rao Lab at Example University for support.")
    d.add_heading("Conflicts of Interest", 1)
    d.add_paragraph("The authors declare no conflict of interest.")
    out = os.path.join(tmp, "m.docx")
    d.save(out)
    return out


def test_requirements_status_detail_and_evidence():
    reqs = {r.key: r for r in sp.extract_requirements(GUIDE)}
    assert reqs["cover_letter"].status == sp.REQUIRED
    assert reqs["graphical_abstract"].status == sp.OPTIONAL
    assert reqs["highlights"].detail == "3–5 bullet points, max 85 characters each"
    assert reqs["reviewers"].detail == "3 reviewers"
    assert "TIFF" in reqs["figures_separate"].detail and "300 dpi" in reqs["figures_separate"].detail
    assert reqs["word_limit"].detail == "5,000 words"
    assert reqs["ethics"].status == sp.NOT_FOUND
    assert "cover letter" in reqs["cover_letter"].evidence


def test_pack_contents_and_honesty():
    with tempfile.TemporaryDirectory() as tmp:
        ms = ingest.ingest(_docx(tmp), image_dir=os.path.join(tmp, "img"))
        g = sp.Guidelines(journal="Renewable Energy", publisher="Elsevier",
                          pages=["https://example.org/guide"], text=GUIDE)
        zbytes, reqs = sp.build_pack(ms, profiles.get("elsevier_generic"), g,
                                     include_latex=False)
        z = zipfile.ZipFile(io.BytesIO(zbytes))
        names = [n.split("/", 1)[1] for n in z.namelist()]
        for want in ("01_Manuscript.docx", "02_Title_page.docx", "03_Cover_letter.docx",
                     "04_Manuscript_blinded.docx", "05_Declarations.docx",
                     "06_Highlights.docx", "08_Suggested_reviewers.csv",
                     "Figures/Figure_1.tif", "00_Submission_checklist.md"):
            assert want in names, want
        assert "07_Tables.docx" not in names                  # no tables in the paper

        from docx import Document
        read = lambda n: "\n".join(p.text for p in Document(io.BytesIO(
            z.read(next(x for x in z.namelist() if x.endswith(n))))).paragraphs)
        blind = read("04_Manuscript_blinded.docx")
        assert "Asha Rao" not in blind and "Rao Lab" not in blind
        assert "Asha Rao" in read("01_Manuscript.docx")
        decl = read("05_Declarations.docx")
        assert "The authors declare no conflict of interest." in decl   # from manuscript
        assert "[Funder name" in decl                                   # placeholder, not invented
        letter = read("03_Cover_letter.docx")
        assert "Renewable Energy" in letter and "[confirm]" in letter
        hl = read("06_Highlights.docx")
        assert "It improves detection accuracy by 12 percent." in hl

        rev = z.read(next(x for x in z.namelist() if x.endswith(".csv") and "review" in x))
        assert rev.decode().count("\n") == 4                            # header + 3 rows
        md = z.read(next(x for x in z.namelist() if x.endswith(".md"))).decode()
        assert "Within limit" in md and "https://example.org/guide" in md


def test_gather_follows_publisher_guide_path():
    pages = {
        "https://www.sciencedirect.com/journal/renewable-energy/publish/guide-for-authors":
            "<html><title>Guide for authors - Renewable Energy</title><body>" + GUIDE + "</body></html>",
        "https://www.sciencedirect.com/journal/renewable-energy":
            "<html><body>Home</body></html>",
    }

    def fake_fetch(u):
        if u not in pages:
            raise OSError("404")
        return u, pages[u]
    with mock.patch("core.discover.fetch", side_effect=fake_fetch):
        g = sp.gather_guidelines(url="https://www.sciencedirect.com/journal/renewable-energy")
    assert g.pages[0].endswith("/publish/guide-for-authors")
    assert "Highlights are mandatory" in g.text


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
