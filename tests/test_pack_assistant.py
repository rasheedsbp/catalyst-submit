"""Fill-in assistant — guided mode, offline (no API key).

Run:  python tests/test_pack_assistant.py
"""
import io
import os
import sys
import tempfile
import zipfile
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from core import ingest, profiles, submission_pack as sp, pack_assistant as pa  # noqa: E402
import test_submission_pack as tsp  # noqa: E402


def _pack():
    tmp = tempfile.mkdtemp()
    ms = ingest.ingest(tsp._docx(tmp), image_dir=os.path.join(tmp, "img"))
    g = sp.Guidelines(journal="Renewable Energy", text=tsp.GUIDE)
    z, reqs = sp.build_pack(ms, profiles.get("elsevier_generic"), g, include_latex=False)
    return ms, reqs, z


def _read(z, suffix):
    from docx import Document
    zz = zipfile.ZipFile(io.BytesIO(z))
    return Document(io.BytesIO(zz.read(next(n for n in zz.namelist() if n.endswith(suffix)))))


def test_scan_groups_same_answer_across_files():
    _ms, _r, z = _pack()
    fields = pa.scan_pack(z)
    email = next(f for f in fields if f.key == "email:Asha Rao")
    assert {loc[0].split("/")[-1] for loc in email.locations} >= {"02_Title_page.docx",
                                                                   "03_Cover_letter.docx"}
    assert any(f.ptype == "credit" for f in fields)
    assert sum(1 for f in fields if f.ptype == "reviewer") == 3


def test_validation():
    f = pa.Field("k", "orcid", "q")
    assert pa.normalise(f, "0000000218250097")[0] == "0000-0002-1825-0097"
    assert pa.normalise(f, "0000-0002-1825-0098")[1]            # bad checksum
    assert pa.normalise(pa.Field("k", "email", "q"), "nope")[1]
    assert pa.normalise(pa.Field("k", "funding", "q"), "none")[0].startswith("This research did not")
    assert pa.normalise(pa.Field("k", "confirm", "q"), "maybe")[1]


def test_conversation_fills_files_offline():
    ms, reqs, z = _pack()
    with mock.patch.object(pa, "llm_reply", return_value=None):   # no key, no network
        s = pa.Session(pa.scan_pack(z), pa.context_for(ms, reqs, "Renewable Energy"))
        assert "CRediT" in s.handle("what is CRediT?")
        s.handle("yes")                                   # confirm
        s.handle("Research Article")                      # article type
        s.handle("It suits the journal's wind-energy focus.")
        s.handle("Dept of EE, NIT Warangal, India")       # affiliation Asha
        s.handle("skip")                                  # affiliation Vikram
        s.handle("NIT Warangal 506004, India")            # postal address
        assert "⚠️" in s.handle("not-an-email")
        s.handle("asha@nitw.ac.in")
    z2, n = pa.apply_answers(z, s.fields)
    title = "\n".join(p.text for p in _read(z2, "02_Title_page.docx").paragraphs)
    letter = "\n".join(p.text for p in _read(z2, "03_Cover_letter.docx").paragraphs)
    assert "Dept of EE, NIT Warangal, India" in title and "asha@nitw.ac.in" in title
    assert "asha@nitw.ac.in" in letter and "[confirm]" not in letter
    assert "(article type: Research Article)" in letter
    assert "[department, institution, city, country]" in title      # skipped stays visible


def test_credit_and_reviewers_written():
    ms, reqs, z = _pack()
    fields = pa.scan_pack(z)
    for f in fields:
        if f.key == "credit:Asha Rao":
            f.answer = "1,6"
        if f.key == "reviewer:1":
            f.answer = "Dr. A. Kumar; IIT Madras; India; ak@iitm.ac.in; condition monitoring"
    z2, _ = pa.apply_answers(z, fields)
    t = _read(z2, "05_Declarations.docx").tables[0]
    assert t.rows[1].cells[1].text == "X" and t.rows[6].cells[1].text == "X"
    assert t.rows[2].cells[1].text == ""
    csv_txt = zipfile.ZipFile(io.BytesIO(z2)).read(
        next(n for n in zipfile.ZipFile(io.BytesIO(z2)).namelist() if n.endswith(".csv"))).decode()
    assert "Dr. A. Kumar,IIT Madras,India,ak@iitm.ac.in,,condition monitoring" in csv_txt


def test_llm_picks_live_model_and_falls_back():
    class R:
        def __init__(self, ok, js, code=200):
            self.ok, self._js, self.status_code = ok, js, code

        def json(self):
            return self._js
    pa._MODEL_CACHE.clear()
    with mock.patch.dict(os.environ, {"GROQ_API_KEY": "x"}, clear=False), \
         mock.patch("requests.get", return_value=R(True, {"data": [
             {"id": "whisper-large-v3"}, {"id": "llama-3.3-70b-versatile"}]})), \
         mock.patch("requests.post", return_value=R(True, {"choices": [
             {"message": {"content": "Draft text"}}]})) as post:
        assert pa.llm_reply("q", "ctx") == "Draft text"
        assert post.call_args.kwargs["json"]["model"] == "llama-3.3-70b-versatile"
    with mock.patch.dict(os.environ, {"GROQ_API_KEY": "x"}, clear=False), \
         mock.patch("requests.post", side_effect=OSError("down")):
        assert pa.llm_reply("q", "ctx") is None            # guided mode takes over


def test_key_name_is_case_insensitive():
    with mock.patch.dict(os.environ, {"Groq_API_Key": " gsk_abc "}, clear=False):
        os.environ.pop("GROQ_API_KEY", None)
        assert pa._secret("GROQ_API_KEY") == "gsk_abc"
        assert pa.llm_status() == "Groq"


def test_diagnose_reports_real_error_and_fallback_models():
    class R:
        def __init__(self, ok, js=None, code=200, text=""):
            self.ok, self._js, self.status_code, self.text = ok, js or {}, code, text

        def json(self):
            return self._js
    pa._MODEL_CACHE.clear()
    env = {"GROQ_API_KEY": "gsk_testkey1234"}
    # model list unreachable -> fallback names are still tried; 401 is reported
    with mock.patch.dict(os.environ, env, clear=False), \
         mock.patch("requests.get", side_effect=OSError("blocked")), \
         mock.patch("requests.post", return_value=R(False, code=401,
                                                     text='{"error":"Invalid API Key"}')):
        report = "\n".join(pa.diagnose())
    assert "key found" in report and "llama-3.3-70b-versatile" in report
    assert "401" in report and "Invalid API Key" in report
    # first model decommissioned (400) -> next model works
    calls = []

    def post(url, **kw):
        calls.append(kw["json"]["model"])
        if len(calls) == 1:
            return R(False, code=400, text="model decommissioned")
        return R(True, {"choices": [{"message": {"content": "OK"}}]})
    with mock.patch.dict(os.environ, env, clear=False), \
         mock.patch("requests.get", return_value=R(True, {"data": []})), \
         mock.patch("requests.post", side_effect=post):
        assert pa.llm_reply("q", "c") == "OK" and len(calls) == 2


def test_manuscript_default_offered_for_name():
    s = pa.Session([pa.Field("corr_name:", "corr_name", "Corresponding author's full name:",
                             locations=[("p/03_Cover_letter.docx", "para", 1, "[name]")])],
                   known={"authors": ["Asha Rao", "Vikram Singh"], "corresponding": None})
    assert "Asha Rao" in s.prompt()
    with mock.patch.object(pa, "llm_reply", return_value=None):
        assert "Asha Rao" in s.handle("suggest")
    s.handle("ok")
    assert s.fields[0].answer == "Asha Rao"


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
