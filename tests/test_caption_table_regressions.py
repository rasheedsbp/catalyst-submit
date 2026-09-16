"""Regressions for the caption / figure / table / class problems found on the
GWSD manuscript runs (Springer, Elsevier and SAGE outputs).

Run:  python -m pytest tests/ -q     (or: python tests/test_caption_table_regressions.py)
"""
import io
import os
import sys
import tempfile
import zipfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core import ingest, inline_math, render_latex, profiles, discover  # noqa: E402


def _png(path):
    from PIL import Image
    Image.new("RGB", (40, 30), (200, 30, 30)).save(path)


def _sample_docx(tmp):
    """Mimics the failing file: unlabelled caption styles, body sentences that
    start with "Table 2 ..."/"Figure 2 ...", unstyled subsection headings."""
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    d = Document()
    for name in ("Table Caption", "Image Caption"):
        d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    d.add_paragraph("A Test Paper Title")
    d.add_paragraph("Jane Doe, John Roe")
    d.add_paragraph("Abstract: This is the abstract of the paper.")
    d.add_paragraph("Keywords: one, two, three")
    d.add_heading("Methods", level=1)
    d.add_paragraph("Simulation design")
    d.add_paragraph("A Monte Carlo simulation assessed the estimating methods "
                    "across many sample sizes and censoring levels in total.")
    d.add_heading("Results", level=1)
    d.add_paragraph("Table 1 compares estimation error across methods and sizes.")
    d.add_paragraph("Parameter-estimation RMSE by method", style="Table Caption")
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["n", "Method", "RMSE"], ["50", "MLE", "27.44"],
                             ["50", "NN", "0.37"]]):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    imgs = []
    for k in (1, 2):
        pth = os.path.join(tmp, f"img{k}.png")
        _png(pth)
        imgs.append(pth)
        d.add_paragraph(f"Figure {k} displays the result of experiment {k} in detail.")
        d.add_picture(pth)
        d.add_paragraph(f"Caption of figure {k}", style="Image Caption")
    d.add_heading("Conclusion", level=1)
    d.add_paragraph("Done.")
    out = os.path.join(tmp, "sample.docx")
    d.save(out)
    return out


def test_body_mentions_are_not_captions_and_figures_do_not_shift():
    with tempfile.TemporaryDirectory() as tmp:
        ms = ingest.ingest(_sample_docx(tmp))
        assert [f.caption for f in ms.figures] == ["Caption of figure 1",
                                                   "Caption of figure 2"]
        assert all(f.path for f in ms.figures)            # no "artwork missing"
        assert ms.tables[0].caption == "Parameter-estimation RMSE by method"
        assert ms.tables[0].rows[1] == ["50", "MLE", "27.44"]
        body = [b.text for s in ms.sections for b in s.blocks if b.kind == "para"]
        assert any(t.startswith("Table 1 compares") for t in body)
        assert any(t.startswith("Figure 2 displays") for t in body)


def test_unstyled_subsection_heading_and_empty_parent_kept():
    with tempfile.TemporaryDirectory() as tmp:
        ms = ingest.ingest(_sample_docx(tmp))
        titles = [(s.level, s.title) for s in ms.sections]
        assert (1, "Methods") in titles
        assert (2, "Simulation design") in titles


def test_inline_math_tokens_render_per_target():
    s = "RMSE(" + inline_math.wrap(r"\alpha") + ")"
    assert inline_math.to_latex(s, render_latex.esc) == r"RMSE($\alpha$)"
    assert inline_math.to_plain(s) == "RMSE(α)"
    assert "allowbreak" not in render_latex._cell_tex(inline_math.wrap(r"c=0,\alpha=1"))


def test_typed_equation_number_is_split_off():
    assert ingest._split_typed_number(r"C=\frac{a}{b} \left( 25 \right)") == (r"C=\frac{a}{b}", "25")
    assert ingest._split_typed_number("a = b(2)") == ("a = b(2)", None)


def test_sage_kit_uses_sagej_not_elsarticle():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("sage/sagej.cls", "\\ProvidesClass{sagej}\n\\def\\corrauth#1{}")
    prof, _ = discover.profile_from_template("sage latex template 4.zip", buf.getvalue())
    assert prof["latex"]["documentclass"].endswith("{sagej}")
    assert prof["latex"]["author_mode"] == "sage"
    assert prof["publisher"] == "SAGE"
    from core.model import Manuscript
    zbytes = render_latex.build_zip(Manuscript(title="T"), prof,
                                    extra_files={"sagej.cls": b"%"})
    z = zipfile.ZipFile(io.BytesIO(zbytes))
    readme = z.read("README.txt").decode()
    assert "sagej.cls" in readme and "elsarticle" not in readme
    assert "{sagej}" in z.read("main.tex").decode().splitlines()[0]




def _mention_docx(tmp):
    from docx import Document
    d = Document()
    d.add_paragraph("Mention Test Title")
    d.add_paragraph("Jane Doe")
    d.add_paragraph("Abstract: short abstract here.")
    d.add_paragraph("Keywords: a, b")
    d.add_heading("Results", level=1)
    d.add_paragraph("Table 1: Dataset summary")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Name", "Size"
    t.cell(1, 0).text, t.cell(1, 1).text = "A", "10"
    # mentions of table numbers in running text — none of these is a table
    d.add_paragraph("Table 3. The results in this table are strong and they clearly "
                    "show that the method works. Later sections discuss it further.")
    d.add_paragraph("Table 4 The proposed method is compared with all baselines here.")
    d.add_paragraph("TABLE 5")
    d.add_paragraph("Results remain stable across folds")
    d.add_paragraph("Tables 2 and 3 are discussed in the next section of the paper.")
    d.add_paragraph("Table 2: Accuracy by model")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Model", "Acc"
    t.cell(1, 0).text, t.cell(1, 1).text = "X", "0.9"
    d.add_heading("Conclusion", level=1)
    d.add_paragraph("Done.")
    out = os.path.join(tmp, "mentions.docx")
    d.save(out)
    return out


def test_table_mentions_in_text_do_not_create_extra_tables():
    with tempfile.TemporaryDirectory() as tmp:
        ms = ingest.ingest(_mention_docx(tmp))
        assert [(t.label, t.caption) for t in ms.tables] == [
            ("tab1", "Dataset summary"), ("tab2", "Accuracy by model")]
        body = " ".join(b.text for s in ms.sections for b in s.blocks if b.kind == "para")
        assert "Table 3. The results" in body
        assert "Table 4 The proposed" in body
        assert "Results remain stable across folds" in body


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_"):
            fn()
            print("ok", name)
