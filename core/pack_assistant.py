"""Fill-in assistant for a generated submission package.

Walks the author through every [placeholder] and [confirm] left in the package
files, one question at a time, writes the answers straight into the Word files
and the reviewer sheet, and returns an updated zip.

It is free in both modes:

  * Guided mode (always on, no key, no internet): the questions, examples,
    validation (emails, ORCID iDs, highlight length) and a built-in answer bank
    for common submission questions ("what is a data availability statement?").
  * AI mode (optional): if a free API key is configured — Groq, Google Gemini or
    OpenRouter, all of which have no-card free tiers — free-form questions and
    "suggest" requests are answered by a hosted model. Model names change often,
    so the model is picked from the provider's live model list, and any failure
    falls back to guided mode instead of breaking the chat.

Everything runs as plain HTTPS calls, so it works on Streamlit Cloud or any VPS.
"""
from __future__ import annotations

import csv
import io
import os
import re
import zipfile
from dataclasses import dataclass, field

# --------------------------------------------------------------------------- #
# fields
# --------------------------------------------------------------------------- #
PLACEHOLDER = re.compile(r"\[(?!x\])([^\[\]]{2,160})\]")


@dataclass
class Field:
    key: str                 # shared across files: same answer fills every copy
    ptype: str
    question: str
    hint: str = ""
    example: str = ""
    person: str | None = None
    locations: list[tuple] = field(default_factory=list)   # (file, kind, loc, placeholder)
    answer: str | None = None
    skipped: bool = False


PTYPES = [
    # ptype, match on placeholder text, question, hint, example
    ("confirm", r"^confirm$",
     "Please confirm this statement is true for your submission (yes / no):",
     "Answer yes to keep it. Answer no and it is marked [NEEDS REVISION] so you can "
     "edit it by hand.", "yes"),
    ("shorten", r"^shorten",
     "This highlight is too long. Type a shorter version:",
     "Keep the key result; drop methods detail. Count includes spaces.", ""),
    ("article_type", r"article type|original research article",
     "What article type are you submitting?",
     "Use the journal's own label from its submission system.",
     "Original Research Article"),
    ("scope", r"suits the journal",
     "In one or two sentences, why does this work suit this journal's scope and readers?",
     "Name the journal's topic your paper advances and what readers gain.",
     "The study advances data-driven fault diagnosis, a core topic of the journal, "
     "and gives practitioners a deployable, low-cost detection model."),
    ("affiliation", r"department|institution|affiliation",
     "Affiliation for {person}:", "Department, institution, city, country.",
     "Department of Computer Science, Osmania University, Hyderabad, India"),
    ("address", r"postal address",
     "Full postal address for the corresponding author {person}:",
     "Include postcode and country.", "Department of CSE, XYZ University, Hyderabad 500007, India"),
    ("email", r"^email$", "Email address for {person}:",
     "Use an institutional address where possible.", "name@university.edu"),
    ("orcid", r"0000-0000", "ORCID iD for {person} (or 'skip'):",
     "16 digits in the form 0000-0002-1825-0097. Free at orcid.org.", "0000-0002-1825-0097"),
    ("funding", r"funder name",
     "Funding: who funded the work and the grant number? (or: 'none')",
     "Name each funder exactly and give grant numbers.",
     "This work was supported by the Science and Engineering Research Board (SERB), "
     "India, under Grant No. CRG/2023/001234."),
    ("data", r"data can be accessed",
     "Data availability: where can readers get the data? (or: 'on request')",
     "Give a repository link/DOI, or state it is available on request.",
     "The dataset is publicly available at https://doi.org/10.5281/zenodo.1234567."),
    ("ethics", r"committee name",
     "Ethics: approval body and number, or 'not applicable':",
     "Needed for human or animal data.", "Not applicable."),
    ("ai", r"ai tool",
     "Did you use any generative-AI tool while writing? Describe it, or 'none':",
     "Most journals require this even when the answer is none.",
     "During the preparation of this work the authors used Grammarly to improve "
     "language. The authors reviewed and edited the content and take full "
     "responsibility for the publication."),
    ("keywords", r"^keywords$", "Keywords (comma-separated):", "", "fault detection, wind energy"),
    ("title", r"^title$", "Manuscript title:", "", ""),
    ("journal", r"journal name", "Target journal name:", "", ""),
    ("corr_name", r"^name$|corresponding author", "Corresponding author's full name:", "", ""),
]

EXPAND = {
    "funding": {"none": "This research did not receive any specific grant from funding "
                        "agencies in the public, commercial, or not-for-profit sectors."},
    "data": {"on request": "Data will be made available on reasonable request.",
             "request": "Data will be made available on reasonable request."},
    "ethics": {"not applicable": "Not applicable.", "na": "Not applicable.", "n/a": "Not applicable."},
    "ai": {"none": "No generative AI or AI-assisted technologies were used in the preparation "
                   "of this manuscript.", "no": "No generative AI or AI-assisted technologies "
                                                "were used in the preparation of this manuscript."},
}


def _classify(ph: str) -> tuple:
    for row in PTYPES:
        if re.search(row[1], ph.strip(), re.I):
            return row
    return ("other", "", f"What should replace “[{ph}]”?", "", "")


def _person(context: str, ph_start: int) -> str | None:
    before = context[:ph_start]
    m = re.search(r"(?:^|\n)\s*\d+\.\s*(.+?)\s+—", before)
    if m:
        return m.group(1).strip()
    line = before.split("\n")[-1]
    m = re.match(r"\s*([A-Z][\w.'\- ]{2,60}?)\s*,", line)
    if m:
        return m.group(1).strip()
    lines = [x for x in context[:ph_start].split("\n") if x.strip() and "[" not in x]
    if lines and re.match(r"^[A-Z][\w.'\-]+(?: [A-Z][\w.'\-]+){0,4}$", lines[-1].strip()):
        return lines[-1].strip()
    return None


def scan_pack(zip_bytes: bytes) -> list[Field]:
    """Every placeholder in the package, grouped so one answer fills all copies."""
    from docx import Document
    fields: dict[str, Field] = {}
    zf = zipfile.ZipFile(io.BytesIO(zip_bytes))
    for name in sorted(zf.namelist()):
        if not name.lower().endswith(".docx") or "/Figures/" in name:
            continue
        if re.search(r"01_Manuscript|04_Manuscript", name):
            continue                                   # authors' own text, never templated
        doc = Document(io.BytesIO(zf.read(name)))
        for pi, p in enumerate(doc.paragraphs):
            text = p.text
            for m in PLACEHOLDER.finditer(text):
                ph = m.group(1)
                ptype, _pat, q, hint, ex = _classify(ph)
                person = _person(text, m.start()) if ptype in (
                    "affiliation", "address", "email", "orcid") else None
                if ptype in ("confirm", "shorten", "other"):
                    key = f"{ptype}:{name}:{pi}:{m.start()}"
                else:
                    key = f"{ptype}:{person or ''}"
                if key not in fields:
                    fields[key] = Field(key=key, ptype=ptype,
                                        question=q.format(person=person or "this author"),
                                        hint=hint, example=ex, person=person)
                    if ptype in ("confirm", "shorten"):
                        fields[key].hint = (fields[key].hint + "\n\n> " +
                                            re.sub(r"\[confirm\]\s*|\s*\[shorten[^\]]*\]", "", text).strip()[:400])
                fields[key].locations.append((name, "para", pi, f"[{ph}]"))
        # CRediT table: one question per author
        for ti, t in enumerate(doc.tables):
            if not t.rows or t.rows[0].cells[0].text.strip() != "Role":
                continue
            for ci, cell in enumerate(t.rows[0].cells[1:], 1):
                person = cell.text.strip()
                key = f"credit:{person}"
                fields.setdefault(key, Field(
                    key=key, ptype="credit", person=person,
                    question=f"CRediT roles for {person}? (numbers, e.g. 1,6,13)",
                    hint="\n".join(f"{i}. {r}" for i, r in enumerate(CREDIT_ROLES, 1)),
                    example="1,6,9,13"))
                fields[key].locations.append((name, "credit", (ti, ci), ""))
    for name in zf.namelist():
        if name.endswith("08_Suggested_reviewers.csv"):
            rows = list(csv.reader(io.StringIO(zf.read(name).decode())))
            for ri in range(1, len(rows)):
                if not any(c.strip() for c in rows[ri]):
                    key = f"reviewer:{ri}"
                    fields[key] = Field(
                        key=key, ptype="reviewer",
                        question=f"Suggested reviewer {ri}: name; affiliation; country; "
                                 "email; expertise (separate with ;) — or 'skip'",
                        hint="Pick active researchers in your topic with no co-authorship "
                             "or shared affiliation with you in the last 3–5 years. "
                             "Check their recent papers for emails.",
                        example="Dr. A. Kumar; IIT Madras; India; akumar@iitm.ac.in; "
                                "wind turbine condition monitoring")
                    fields[key].locations.append((name, "csv", ri, ""))
    order = ["confirm", "article_type", "scope", "affiliation", "address", "email", "orcid",
             "funding", "data", "ethics", "ai", "credit", "shorten", "reviewer",
             "keywords", "title", "journal", "corr_name", "other"]
    return sorted(fields.values(), key=lambda f: (order.index(f.ptype)
                                                  if f.ptype in order else 99, f.person or ""))


from .submission_pack import CREDIT_ROLES  # noqa: E402


# --------------------------------------------------------------------------- #
# validation
# --------------------------------------------------------------------------- #
def normalise(f: Field, raw: str) -> tuple[str | None, str | None]:
    """(value, error). value None + error = ask again."""
    a = (raw or "").strip()
    low = a.lower().rstrip(".")
    if f.ptype in EXPAND and low in EXPAND[f.ptype]:
        return EXPAND[f.ptype][low], None
    if f.ptype == "confirm":
        if low in ("y", "yes", "true", "confirm", "confirmed", "ok", "haan", "ha"):
            return "yes", None
        if low in ("n", "no", "false", "nahi"):
            return "no", None
        return None, "Please answer yes or no."
    if f.ptype == "email":
        return (a, None) if re.fullmatch(r"[^@\s]+@[^@\s]+\.[A-Za-z]{2,}", a) else \
            (None, "That doesn't look like an email address — try again.")
    if f.ptype == "orcid":
        digits = re.sub(r"[^0-9Xx]", "", a)
        if len(digits) != 16:
            return None, "An ORCID iD has 16 characters, e.g. 0000-0002-1825-0097."
        if not _orcid_ok(digits):
            return None, "That ORCID iD fails its checksum — please re-check the digits."
        return "-".join(digits[i:i + 4] for i in range(0, 16, 4)).upper(), None
    if f.ptype == "shorten":
        m = re.search(r"shorten: (\d+)", f.locations[0][3]) if f.locations else None
        limit = 85
        if len(a) > limit:
            return None, f"Still {len(a)} characters — it must be {limit} or fewer."
        return a, None
    if f.ptype == "credit":
        nums = [int(x) for x in re.findall(r"\d+", a) if 1 <= int(x) <= len(CREDIT_ROLES)]
        if not nums:
            return None, "Give role numbers from the list, e.g. 1,6,13."
        return ",".join(str(n) for n in sorted(set(nums))), None
    if f.ptype == "reviewer":
        parts = [p.strip() for p in a.split(";")]
        if len(parts) < 2:
            return None, "Separate the details with ; — at least name and affiliation."
        if len(parts) >= 4 and parts[3] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.\w{2,}", parts[3]):
            return None, "The 4th item should be an email address."
        return a, None
    return (a, None) if a else (None, "Please type an answer, or 'skip'.")


def _orcid_ok(d: str) -> bool:
    total = 0
    for ch in d[:15]:
        total = (total + int(ch)) * 2
    chk = (12 - total % 11) % 11
    return d[15].upper() == ("X" if chk == 10 else str(chk))


# --------------------------------------------------------------------------- #
# writing answers back
# --------------------------------------------------------------------------- #
def _replace_in_paragraph(p, placeholder: str, value: str) -> bool:
    for run in p.runs:                               # usual case: inside one run
        if placeholder in run.text:
            run.text = run.text.replace(placeholder, value, 1)
            return True
    full = p.text
    if placeholder not in full:
        return False
    new = full.replace(placeholder, value, 1)
    for i, run in enumerate(p.runs):
        run.text = new if i == 0 else ""
    return True


def apply_answers(zip_bytes: bytes, fields: list[Field]) -> tuple[bytes, int]:
    from docx import Document
    zin = zipfile.ZipFile(io.BytesIO(zip_bytes))
    by_file: dict[str, list[tuple[Field, tuple]]] = {}
    for f in fields:
        if f.answer is None:
            continue
        for loc in f.locations:
            by_file.setdefault(loc[0], []).append((f, loc))
    filled = 0
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in zin.namelist():
            data = zin.read(name)
            jobs = by_file.get(name)
            if jobs and name.endswith(".docx"):
                doc = Document(io.BytesIO(data))
                for f, (_n, kind, loc, ph) in jobs:
                    if kind == "para":
                        p = doc.paragraphs[loc]
                        if f.ptype == "confirm":
                            val = "" if f.answer == "yes" else "[NEEDS REVISION]"
                            ok = _replace_in_paragraph(p, ph + " ", val + (" " if val else "")) \
                                or _replace_in_paragraph(p, ph, val)
                        elif f.ptype == "shorten":
                            ok = True
                            for i, run in enumerate(p.runs):
                                run.text = f.answer if i == 0 else ""
                        else:
                            ok = _replace_in_paragraph(p, ph, f.answer)
                        filled += bool(ok)
                    elif kind == "credit":
                        ti, ci = loc
                        t = doc.tables[ti]
                        roles = {int(x) for x in f.answer.split(",")}
                        for ri, row in enumerate(t.rows[1:], 1):
                            row.cells[ci].text = "X" if ri in roles else ""
                        filled += 1
                buf = io.BytesIO()
                doc.save(buf)
                data = buf.getvalue()
            elif jobs and name.endswith(".csv"):
                rows = list(csv.reader(io.StringIO(data.decode())))
                for f, (_n, _k, ri, _ph) in jobs:
                    parts = [p.strip() for p in f.answer.split(";")]
                    # name; affiliation; country; email; expertise (+ why)
                    row = (parts + [""] * 7)[:7]
                    if len(parts) == 5:
                        row = [parts[0], parts[1], parts[2], parts[3], "", parts[4], ""]
                    rows[ri] = row
                    filled += 1
                buf = io.StringIO()
                csv.writer(buf).writerows(rows)
                data = buf.getvalue().encode()
            zout.writestr(name, data)
    return out.getvalue(), filled


# --------------------------------------------------------------------------- #
# answer bank (guided mode, no key)
# --------------------------------------------------------------------------- #
FAQ = [
    (r"credit|author contribution",
     "**CRediT** is a standard list of 14 contributor roles (Conceptualization, "
     "Methodology, Software, Writing – original draft, …). Each author gets the roles "
     "they actually did; one role can have several authors. Most journals print it as "
     "“Author A: Conceptualization, Methodology. Author B: Software, Validation.”"),
    (r"data availab",
     "A **data availability statement** tells readers where the data behind the results "
     "can be found. Typical forms: a repository link/DOI (Zenodo, Figshare, Mendeley "
     "Data, GitHub); “available on reasonable request”; or “no data were used”. Public "
     "datasets: cite them and give the link."),
    (r"competing|conflict of interest",
     "A **competing-interest declaration** discloses anything that could look like bias: "
     "funding from a company with a stake, patents, consultancy, personal ties. If none, "
     "the standard wording already in your declarations file is correct — just confirm it."),
    (r"highlight",
     "**Highlights** are 3–5 bullets, usually ≤85 characters each including spaces, "
     "stating the main findings in plain words. Start each with the result, not “In this "
     "paper”. Example: “Compact CNN detects turbine faults 12% more accurately.”"),
    (r"orcid",
     "An **ORCID iD** is a free, permanent researcher ID (orcid.org, 2-minute sign-up). "
     "Many journals require it for the corresponding author and encourage it for all."),
    (r"reviewer",
     "**Suggested reviewers** should be active in your exact topic, from different "
     "institutions and countries, with no co-authorship or shared affiliation with you "
     "in the last 3–5 years. Take them from recent papers you cite; use institutional "
     "emails. Editors may not invite them, so avoid friends or collaborators."),
    (r"cover letter",
     "A **cover letter** states the title and article type, summarises the contribution "
     "in 2–3 sentences, explains the fit with the journal's scope, and confirms the paper "
     "is original and not under review elsewhere. Keep it to one page."),
    (r"ethic",
     "**Ethics approval** is needed for studies on humans, human data or animals: give "
     "the committee name and approval number, and say informed consent was obtained. "
     "For purely computational work on public data, “Not applicable” is usual."),
    (r"blind|anonym",
     "For **double-blind review**, the blinded manuscript must not reveal authors: names, "
     "affiliations, acknowledgements, funding grant numbers, and self-citations written "
     "as “our previous work [12]” (write “Kumar et al. [12]” instead)."),
    (r"graphical abstract",
     "A **graphical abstract** is one image summarising the paper (often ~1328×531 px "
     "minimum for Elsevier, readable at small size). Tools: PowerPoint, BioRender, "
     "Canva. It cannot be generated automatically."),
    (r"article type",
     "Use the journal's own label from its submission system — most research papers are "
     "“Research Article” / “Original Research Article”; reviews are “Review Article”."),
    (r"\bai\b|generative|chatgpt",
     "Many publishers require a **generative-AI declaration**: name the tool, what it was "
     "used for (e.g. language editing), and state the authors reviewed the output and take "
     "responsibility. AI tools cannot be listed as authors."),
]


def faq_answer(text: str) -> str | None:
    for pat, ans in FAQ:
        if re.search(pat, text, re.I):
            return ans
    return None


# --------------------------------------------------------------------------- #
# optional free LLM
# --------------------------------------------------------------------------- #
PROVIDERS = [
    # name, env/secret key, OpenAI-compatible base URL, preferred model substrings
    ("Groq", "GROQ_API_KEY", "https://api.groq.com/openai/v1",
     ["llama-3.3-70b", "llama-4", "gpt-oss-120b", "qwen", "llama"]),
    ("Gemini", "GEMINI_API_KEY", "https://generativelanguage.googleapis.com/v1beta/openai",
     ["gemini-2.5-flash", "gemini-flash", "gemini-2", "flash"]),
    ("OpenRouter", "OPENROUTER_API_KEY", "https://openrouter.ai/api/v1",
     [":free"]),
]
_MODEL_CACHE: dict[str, str] = {}


def _secret(name: str) -> str | None:
    """Key from env or Streamlit secrets, whatever the capitalisation.

    Secret names are case-sensitive, so "Groq_API_Key" would never match
    "GROQ_API_KEY". Match ignoring case (and '-' vs '_'), and also look one
    level down, for keys kept under a [section] in secrets.toml.
    """
    want = name.lower().replace("-", "_")
    norm = lambda k: str(k).strip().lower().replace("-", "_")
    for k, v in os.environ.items():
        if norm(k) == want and v.strip():
            return v.strip()
    try:
        import streamlit as st
        sec = st.secrets
        for k in list(sec.keys()):
            v = sec[k]
            if norm(k) == want and isinstance(v, str) and v.strip():
                return v.strip()
            if hasattr(v, "keys"):
                for k2 in list(v.keys()):
                    if norm(k2) == want and isinstance(v[k2], str) and v[k2].strip():
                        return v[k2].strip()
    except Exception:  # noqa: BLE001
        pass
    return None


def llm_status() -> str | None:
    for name, key, _b, _m in PROVIDERS:
        if _secret(key):
            return name
    return None


LAST_ERROR: dict[str, str] = {}
# used when the model list cannot be read; tried in order
FALLBACK_MODELS = {
    "Groq": ["llama-3.3-70b-versatile", "openai/gpt-oss-120b", "llama-3.1-8b-instant",
             "meta-llama/llama-4-scout-17b-16e-instruct"],
    "Gemini": ["gemini-2.5-flash", "gemini-2.0-flash", "gemini-flash-latest"],
    "OpenRouter": ["meta-llama/llama-3.3-70b-instruct:free", "openai/gpt-oss-20b:free"],
}


def _models(name: str, base: str, key: str, prefs: list[str]) -> list[str]:
    """Candidate models: preferred matches from the live list, then fallbacks."""
    import requests
    live: list[str] = []
    try:
        r = requests.get(base + "/models", headers={"Authorization": f"Bearer {key}"},
                         timeout=15)
        if r.ok:
            live = [str(m.get("id", "")).replace("models/", "")
                    for m in r.json().get("data", [])]
        else:
            LAST_ERROR[name] = f"model list: HTTP {r.status_code} {r.text[:160]}"
    except Exception as exc:  # noqa: BLE001
        LAST_ERROR[name] = f"model list: {type(exc).__name__}: {exc}"
    live = [i for i in live if i and not re.search(
        r"whisper|tts|embed|guard|vision|image|audio|playai|compound|prompt", i, re.I)]
    picked = []
    for pref in prefs:
        picked += sorted((i for i in live if pref in i and i not in picked), reverse=True)
    cands = ([_MODEL_CACHE[base]] if base in _MODEL_CACHE else []) + picked + \
        [m for m in FALLBACK_MODELS.get(name, []) if not live or m in live]
    return list(dict.fromkeys(cands))[:5]


SYSTEM = ("You help an academic author complete journal-submission documents. Be concise "
          "(under 120 words), practical and accurate. Never invent facts about the "
          "author's study, funders, approvals or data — if you draft text, mark unknowns "
          "in [brackets] for the author to fill. Use the context given.")


def llm_reply(question: str, context: str) -> str | None:
    import requests
    for name, keyname, base, prefs in PROVIDERS:
        key = _secret(keyname)
        if not key:
            continue
        for model in _models(name, base, key, prefs):
            try:
                r = requests.post(base + "/chat/completions", timeout=45,
                                  headers={"Authorization": f"Bearer {key}",
                                           "Content-Type": "application/json"},
                                  json={"model": model, "temperature": 0.3,
                                        "max_tokens": 400,
                                        "messages": [
                                            {"role": "system", "content": SYSTEM},
                                            {"role": "user",
                                             "content": f"Context:\n{context[:3500]}\n\n"
                                                        f"Question: {question}"}]})
            except Exception as exc:  # noqa: BLE001
                LAST_ERROR[name] = f"{model}: {type(exc).__name__}: {exc}"
                continue
            if r.ok:
                try:
                    txt = r.json()["choices"][0]["message"]["content"].strip()
                except Exception:  # noqa: BLE001
                    txt = ""
                if txt:
                    _MODEL_CACHE[base] = model
                    LAST_ERROR.pop(name, None)
                    return txt
            LAST_ERROR[name] = f"{model}: HTTP {r.status_code} {r.text[:200]}"
            _MODEL_CACHE.pop(base, None)
            if r.status_code in (401, 403):
                break                      # bad key: other models won't help
            if r.status_code == 429:
                break                      # rate limited: try the next provider
    return None


def diagnose() -> list[str]:
    """Human-readable connection report for the page's 'Test AI' button."""
    out = []
    found = False
    for name, keyname, base, prefs in PROVIDERS:
        key = _secret(keyname)
        if not key:
            continue
        found = True
        out.append(f"{name}: key found ({key[:4]}…{key[-4:]}, {len(key)} chars)")
        if name == "Groq" and not key.startswith("gsk_"):
            out.append("  ⚠️ Groq keys normally start with 'gsk_' — check you copied the whole key.")
        cands = _models(name, base, key, prefs)
        out.append(f"  models to try: {', '.join(cands) or 'none'}")
        reply = None
        LAST_ERROR.pop(name, None)
        if cands:
            reply = llm_reply("Reply with the single word OK.", "connection test")
        out.append("  ✅ working — replied: " + reply[:40] if reply else
                   "  ❌ " + LAST_ERROR.get(name, "no reply"))
        if reply:
            break
    if not found:
        out.append("No API key found. In Streamlit Cloud: ⋮ → Settings → Secrets, add a line "
                   'like  GROQ_API_KEY = "gsk_..."  then Save and reboot the app.')
    return out


# --------------------------------------------------------------------------- #
# the conversation
# --------------------------------------------------------------------------- #
HELP = ("Type your answer, or: **skip** · **back** · **example** · **suggest** · "
        "**status** · **done**. Ask any question (e.g. *what is CRediT?*) at any time.")


@dataclass
class Session:
    fields: list[Field]
    context: str = ""
    i: int = 0
    known: dict = field(default_factory=dict)     # facts from the manuscript

    def default_for(self, f: Field) -> str | None:
        k = self.known
        if f.ptype == "corr_name":
            return k.get("corresponding") or (k.get("authors") or [None])[0]
        if f.ptype in ("title", "journal", "keywords"):
            return k.get(f.ptype) or None
        if f.ptype == "article_type":
            return "Original Research Article"
        return None

    def current(self) -> Field | None:
        while self.i < len(self.fields) and (self.fields[self.i].answer is not None
                                             or self.fields[self.i].skipped):
            self.i += 1
        return self.fields[self.i] if self.i < len(self.fields) else None

    def progress(self) -> tuple[int, int]:
        done = sum(1 for f in self.fields if f.answer is not None)
        return done, len(self.fields)

    def prompt(self) -> str:
        f = self.current()
        if f is None:
            d, n = self.progress()
            return (f"✅ All done — {d} of {n} items filled. Download the updated package "
                    "below. Anything you skipped still shows its [placeholder].")
        d, n = self.progress()
        files = sorted({loc[0].split("/")[-1] for loc in f.locations})
        msg = f"**({d + 1}/{n}) {f.question}**\n\n*Goes into: {', '.join(files)}*"
        dflt = self.default_for(f)
        if dflt:
            msg += f"\n\nFrom your manuscript: **{dflt}** — type **ok** to use it."
        if f.hint:
            msg += f"\n\n{f.hint}"
        return msg

    def handle(self, text: str) -> str:
        t = (text or "").strip()
        low = t.lower()
        f = self.current()
        if low in ("help", "?"):
            return HELP
        if low == "status":
            d, n = self.progress()
            left = [x.question.split("?")[0][:60] for x in self.fields
                    if x.answer is None and not x.skipped][:8]
            return f"{d}/{n} filled.\n\nStill open:\n" + "\n".join(f"- {q}" for q in left)
        if low == "done":
            self.i = len(self.fields)
            return self.prompt()
        if low == "back":
            prev = [k for k in range(min(self.i, len(self.fields)) - 1, -1, -1)
                    if self.fields[k].answer is not None or self.fields[k].skipped]
            if prev:
                self.i = prev[0]
                self.fields[self.i].answer, self.fields[self.i].skipped = None, False
            return self.prompt()
        if f is None:
            return self._question(t) or self.prompt()
        if low in ("ok", "okay", "accept", "use it", "yes please") and self.default_for(f) \
                and f.ptype != "confirm":
            f.answer = self.default_for(f)
            return "Saved ✔\n\n" + self.prompt()
        if low == "skip":
            f.skipped = True
            return "Skipped.\n\n" + self.prompt()
        if low == "example":
            return (f"Example: *{f.example}*" if f.example else "No example for this one.") \
                + "\n\n" + self.prompt()
        if low == "suggest":
            dflt = self.default_for(f)
            if dflt:
                return (f"Suggestion from your manuscript: **{dflt}** — type **ok** to use "
                        "it, or type a different answer.")
            ai = llm_reply(f"Draft a suitable answer for: {f.question}", self.context)
            if ai:
                return f"Suggested draft (edit before using):\n\n{ai}\n\nPaste or type your final answer."
            why = ""
            if llm_status():
                err = next(iter(LAST_ERROR.values()), "no reply")
                why = f"\n\n(AI is configured but did not answer: {err[:160]}. Use **Test AI connection** below.)"
            return ((f"Here's a model answer to adapt: *{f.example}*" if f.example
                     else "No automatic suggestion for this one — please type it.")
                    + why + "\n\n" + self.prompt())
        if self._looks_like_question(t):
            return (self._question(t) or "I don't have an answer for that offline.") \
                + "\n\n" + self.prompt()
        value, err = normalise(f, t)
        if err:
            return f"⚠️ {err}\n\n" + self.prompt()
        f.answer = value
        return "Saved ✔\n\n" + self.prompt()

    @staticmethod
    def _looks_like_question(t: str) -> bool:
        return t.endswith("?") or bool(re.match(
            r"^(what|how|why|which|who|when|where|can|should|do|does|is|are|explain|kya)\b",
            t, re.I))

    def _question(self, t: str) -> str | None:
        return faq_answer(t) or llm_reply(t, self.context)


def known_for(ms, journal: str) -> dict:
    authors = [a.name for a in getattr(ms, "authors", []) or []]
    corr = next((a.name for a in getattr(ms, "authors", []) or [] if a.corresponding), None)
    return {"title": getattr(ms, "title", "") or "", "journal": journal or "",
            "authors": authors, "corresponding": corr,
            "keywords": ", ".join(getattr(ms, "keywords", []) or [])}


def context_for(ms, reqs, journal: str) -> str:
    lines = [f"Journal: {journal}", f"Title: {getattr(ms, 'title', '')}",
             f"Abstract: {getattr(ms, 'abstract', '')[:1200]}",
             "Authors: " + ", ".join(a.name for a in getattr(ms, "authors", []))]
    for r in reqs or []:
        if r.status != "Not mentioned":
            lines.append(f"Requirement: {r.label} ({r.status}) {r.detail} — {r.evidence[:160]}")
    return "\n".join(lines)
